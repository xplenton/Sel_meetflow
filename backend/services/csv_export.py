"""CSV + XLSX export helpers.

Extracted from `routes/resources/_common.py` in iter 349. Public API:

    `csv_line(values)`              → str   — one RFC-4180-quoted CSV line
    `build_xlsx(rows, headers, filename)` → bytes — branded XLSX with header row

Legacy underscore-prefixed names `_csv_line` / `_build_xlsx` are kept as
thin shims in `routes/resources/_common.py` for backwards compatibility
with the existing call sites in `routes/resources/invoices.py`.
"""
from __future__ import annotations

import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def csv_line(values: List) -> str:
    """Render a single CSV row with RFC-4180 quoting (comma, quote, newline)."""
    out = []
    for v in values:
        s = "" if v is None else str(v)
        if any(ch in s for ch in [",", "\"", "\n"]):
            s = '"' + s.replace('"', '""') + '"'
        out.append(s)
    return ",".join(out) + "\n"


def build_xlsx(rows: list, headers: list, filename: str) -> bytes:
    """Render a branded XLSX with a styled header row. `filename` is currently
    unused (XLSX has no embedded filename) but kept for API symmetry with
    older call sites."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4A5D4E", end_color="4A5D4E", fill_type="solid")
    for r in rows:
        ws.append(r)
    for col_idx, _ in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
